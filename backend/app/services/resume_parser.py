import os
import pypdf
import docx
from fastapi import HTTPException, status
from app.utils.text_processing import clean_text

class ResumeParserService:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        text_content = []
        try:
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_content.append(extracted)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to read PDF resume: {str(e)}"
            )
        
        full_text = "\n".join(text_content)
        cleaned = clean_text(full_text)
        if not cleaned or len(cleaned.strip()) < 20:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Scanned or image-only resumes cannot be processed. No readable text could be extracted. Please upload a digital PDF or DOCX file with selectable text (no scanned copies or images)."
            )
        return cleaned

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            # Also extract text from table cells if present
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text not in paragraphs:
                            paragraphs.append(cell.text)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to read DOCX resume: {str(e)}"
            )
        
        full_text = "\n".join(paragraphs)
        cleaned = clean_text(full_text)
        if not cleaned or len(cleaned.strip()) < 20:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Scanned or image-only documents cannot be processed. No readable text could be extracted from this DOCX document. Please upload a document with selectable text."
            )
        return cleaned

    @classmethod
    def parse_file(cls, file_path: str, file_type: str) -> str:
        if file_type.upper() == "PDF":
            return cls.extract_text_from_pdf(file_path)
        elif file_type.upper() == "DOCX":
            return cls.extract_text_from_docx(file_path)
        else:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail=f"Unsupported file type: {file_type}"
            )
